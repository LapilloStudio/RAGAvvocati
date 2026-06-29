"""Document text extraction for the professional formats in scope: PDF, DOCX, XLSX.

Returns a list of (page, text) segments. `page` is the 1-based source page where the
format has one (PDF); for DOCX/XLSX it is a logical section index used for citations.
"""

from __future__ import annotations

import io
from dataclasses import dataclass


@dataclass(slots=True)
class ParsedSegment:
    page: int
    text: str


# MIME types we accept at the ingestion boundary.
SUPPORTED_MIME_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}


class UnsupportedDocumentError(ValueError):
    pass


def parse(data: bytes, mime_type: str) -> list[ParsedSegment]:
    kind = SUPPORTED_MIME_TYPES.get(mime_type)
    if kind == "pdf":
        return _parse_pdf(data)
    if kind == "docx":
        return _parse_docx(data)
    if kind == "xlsx":
        return _parse_xlsx(data)
    raise UnsupportedDocumentError(f"Unsupported mime_type: {mime_type!r}")


def _parse_pdf(data: bytes) -> list[ParsedSegment]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    segments: list[ParsedSegment] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            segments.append(ParsedSegment(page=i, text=text))
    return segments


def _parse_docx(data: bytes) -> list[ParsedSegment]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [ParsedSegment(page=1, text=text)] if text.strip() else []


def _parse_xlsx(data: bytes) -> list[ParsedSegment]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    segments: list[ParsedSegment] = []
    for idx, sheet in enumerate(wb.worksheets, start=1):
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            segments.append(ParsedSegment(page=idx, text="\n".join(rows)))
    wb.close()
    return segments
